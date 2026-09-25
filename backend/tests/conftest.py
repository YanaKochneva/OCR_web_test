"""Shared fixtures: isolated data dir, sample documents, synthetic PDFs."""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from PIL import Image

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.schemas.document import (  # noqa: E402
    DocumentJSON,
    DocumentPage,
    DocumentSource,
    DocumentType,
    make_image_block,
    make_table_block,
    make_text_block,
    renumber_blocks,
)

# --------------------------------------------------------------------------- #
#  document builders
# --------------------------------------------------------------------------- #


def build_document(
    document_type: DocumentType,
    *,
    source: DocumentSource = DocumentSource.reference,
    session_id: str = "test-session",
    pages: list[list] | None = None,
    page_numbers: list[int] | None = None,
) -> DocumentJSON:
    """Build a document from a compact description.

    `pages` is a list of pages; each page is a list of blocks given as
    ("text", content) / ("table", rows) / ("image", url) tuples.
    """
    pages = pages if pages is not None else [[("text", "Пример текста.")]]
    numbers = page_numbers or list(range(1, len(pages) + 1))

    doc_pages: list[DocumentPage] = []
    for index, spec in enumerate(pages):
        blocks = []
        for kind, payload in spec:
            if kind == "text":
                blocks.append(make_text_block(payload, source="test"))
            elif kind == "table":
                blocks.append(make_table_block(payload, source="test"))
            elif kind == "image":
                blocks.append(make_image_block(payload, [0, 0, 10, 10]))
            else:  # pragma: no cover - guard against typos in tests
                raise ValueError(f"unknown block kind: {kind}")
        doc_pages.append(
            DocumentPage(
                page_number=numbers[index],
                width=1000,
                height=1400,
                blocks=blocks,
            )
        )

    document = DocumentJSON(
        document_type=document_type,
        source=source,
        session_id=session_id,
        pages=doc_pages,
        parser="test",
    )
    return renumber_blocks(document)


@pytest.fixture
def reference_text_only() -> DocumentJSON:
    return build_document(
        DocumentType.text_only,
        pages=[[("text", "Договор № 42 от 1 января 2026 года.")]],
    )


@pytest.fixture
def table_rows() -> list[list[str]]:
    return [
        ["Наименование", "Кол-во", "Цена"],
        ["Болт М8", "12", "5,50"],
        ["Гайка М8", "24", "2,20"],
    ]


# PyMuPDF's built-in base-14 fonts have no Cyrillic glyphs, so a system TrueType
# font is used whenever the fixture text contains non-ASCII characters.
CYRILLIC_FONT_CANDIDATES = (
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\segoeui.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
)

_CYRILLIC_FONT: str | None = None
_CYRILLIC_FONT_RESOLVED = False


def cyrillic_font() -> str | None:
    """Path to a Cyrillic-capable TTF, or ``None`` when the system has none."""
    global _CYRILLIC_FONT, _CYRILLIC_FONT_RESOLVED
    if not _CYRILLIC_FONT_RESOLVED:
        _CYRILLIC_FONT = next(
            (path for path in CYRILLIC_FONT_CANDIDATES if Path(path).exists()), None
        )
        _CYRILLIC_FONT_RESOLVED = True
    return _CYRILLIC_FONT


def _font_kwargs(text: str) -> dict:
    if all(ord(char) < 128 for char in text):
        return {}
    font = cyrillic_font()
    if font is None:
        pytest.skip("нет TrueType-шрифта с кириллицей для генерации тестового PDF")
    return {"fontname": "cyr", "fontfile": font}


def make_text_pdf(path: Path, lines: list[str], pages: int = 1) -> Path:
    """A digital PDF with a real text layer."""
    import pymupdf

    document = pymupdf.open()
    for page_index in range(pages):
        page = document.new_page()
        y = 72.0
        for line in lines:
            text = f"{line} (стр. {page_index + 1})"
            page.insert_text((72, y), text, fontsize=11, **_font_kwargs(text))
            y += 18
    document.save(str(path))
    document.close()
    return path


def make_scanned_pdf(path: Path, pages: int = 1) -> Path:
    """A PDF whose pages are pure bitmaps - exactly like file B."""
    import pymupdf

    document = pymupdf.open()
    for _ in range(pages):
        page = document.new_page()
        image = Image.new("RGB", (900, 1200), "white")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        page.insert_image(pymupdf.Rect(0, 0, 595, 793), stream=buffer.getvalue())
    document.save(str(path))
    document.close()
    return path


def make_image_pdf(path: Path, image_size: tuple[int, int] = (400, 300)) -> Path:
    """A text PDF that also contains one embedded raster image (a figure)."""
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    caption = "Figure 1 - assembly scheme."
    page.insert_text((72, 72), caption, fontsize=12, **_font_kwargs(caption))
    image = Image.new("RGB", image_size, (30, 120, 200))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    page.insert_image(pymupdf.Rect(72, 120, 372, 345), stream=buffer.getvalue())
    document.save(str(path))
    document.close()
    return path


def make_docx(
    path: Path, paragraphs: list[str], table_rows: list[list[str]] | None = None
) -> Path:
    from docx import Document as DocxDocument

    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    document.save(str(path))
    return path


# --------------------------------------------------------------------------- #
#  app fixtures
# --------------------------------------------------------------------------- #


@pytest.fixture
def test_settings(tmp_path: Path):
    from app.config import Settings

    return Settings(
        glm_mock=True,
        glm_api_key="",
        data_dir=str(tmp_path / "data"),
        render_dpi=72,  # keep synthetic tests fast
        render_max_pages=5,
        text_layer_min_chars=5,
        max_upload_mb=5,
    )


@pytest.fixture
def stores(test_settings):
    from app.services.images import ImageStore
    from app.services.storage import SessionStore

    data_dir = test_settings.data_path
    return SessionStore(data_dir), ImageStore(data_dir)


@pytest.fixture
def api_client(test_settings):
    """FastAPI TestClient wired to an isolated data directory."""
    from fastapi.testclient import TestClient

    from app.api.deps import get_image_store, get_session_store, get_settings_dep
    from app.main import app
    from app.services.images import ImageStore
    from app.services.storage import SessionStore

    data_dir = test_settings.data_path
    store = SessionStore(data_dir)
    images = ImageStore(data_dir)

    app.dependency_overrides[get_settings_dep] = lambda: test_settings
    app.dependency_overrides[get_session_store] = lambda: store
    app.dependency_overrides[get_image_store] = lambda: images

    with TestClient(app) as client:
        client.store = store  # type: ignore[attr-defined]
        yield client

    app.dependency_overrides.clear()


__all__ = [
    "build_document",
    "cyrillic_font",
    "make_docx",
    "make_image_pdf",
    "make_scanned_pdf",
    "make_text_pdf",
]
