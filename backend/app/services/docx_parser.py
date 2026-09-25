"""DOCX reference extraction via python-docx.

Word files have no page raster, so any embedded picture is exported from the
package blob directly; `bbox` is left empty and the UI simply lists such images.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as open_docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.schemas.document import (
    DocumentJSON,
    DocumentPage,
    DocumentSource,
    DocumentType,
    make_image_block,
    make_table_block,
    make_text_block,
    renumber_blocks,
)
from app.services.images import ImageStore
from app.services.layout import sort_document_reading_order

DOCX_MAGIC = b"PK\x03\x04"

IMAGE_CONTENT_TYPES = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/tiff": ".tiff",
    "image/webp": ".webp",
    "image/x-emf": ".emf",
    "image/x-wmf": ".wmf",
}


@dataclass
class DocxExtraction:
    document: DocumentJSON
    warnings: list[str] = field(default_factory=list)


def _iter_block_items(parent) -> list[Paragraph | Table]:
    """Yield paragraphs and tables in document order."""
    if isinstance(parent, DocxDocument):
        parent_element = parent.element.body
    else:  # table cell
        parent_element = parent._tc

    items: list[Paragraph | Table] = []
    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            items.append(Paragraph(child, parent))
        elif child.tag == qn("w:tbl"):
            items.append(Table(child, parent))
    return items


def _table_rows(table: Table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        values: list[str] = []
        for cell in row.cells:
            text = " ".join(
                paragraph.text.strip()
                for paragraph in cell.paragraphs
                if paragraph.text.strip()
            )
            values.append(text)
        if values:
            rows.append(values)
    return rows


def _paragraph_images(paragraph: Paragraph, document: DocxDocument):
    """Return (blob, extension, relation_id) for every picture in a paragraph."""
    found = []
    for blip in paragraph._p.findall(".//" + qn("a:blip")):
        relation_id = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
        if not relation_id:
            continue
        try:
            part = document.part.related_parts[relation_id]
        except KeyError:  # pragma: no cover - broken package
            continue
        extension = IMAGE_CONTENT_TYPES.get(
            getattr(part, "content_type", ""), Path(str(part.partname)).suffix or ".png"
        )
        found.append((part.blob, extension, relation_id))
    return found


def extract_docx_reference(
    path: str | Path,
    *,
    document_type: DocumentType,
    session_id: str,
    store: ImageStore,
) -> DocxExtraction:
    """Build the reference JSON for a DOCX file.

    DOCX flows have no reliable page boundaries, so the whole document is
    reported as a single logical page (`page_number = 1`).  This keeps the JSON
    schema identical to the PDF case.
    """
    entity_set = document_type.entity_set
    warnings: list[str] = []

    try:
        document = open_docx(str(path))
    except Exception as exc:
        raise RuntimeError(f"Не удалось прочитать DOCX: {exc}") from exc

    blocks: list = []
    image_index = 0

    for item in _iter_block_items(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if text and entity_set.text:
                blocks.append(make_text_block(text, source="python-docx"))
            if entity_set.images:
                for blob, extension, relation_id in _paragraph_images(item, document):
                    image_index += 1
                    url = store.save_bytes(
                        session_id,
                        blob,
                        extension=extension,
                        tag=f"docx-{image_index:03d}",
                    )
                    blocks.append(
                        make_image_block(url, source="python-docx")
                    )
        elif isinstance(item, Table):
            if not entity_set.tables:
                continue
            rows = _table_rows(item)
            if rows:
                blocks.append(make_table_block(rows, source="python-docx"))

    document_json = DocumentJSON(
        document_type=document_type,
        source=DocumentSource.reference,
        session_id=session_id,
        pages=[
            DocumentPage(page_number=1, width=None, height=None, blocks=blocks)
        ],
        parser="docx",
    )
    sort_document_reading_order(document_json)
    renumber_blocks(document_json)

    if not entity_set.tables:
        warnings.append("Таблицы DOCX пропущены: тип документа не подразумевает таблицы.")
    warnings.append(
        "DOCX не имеет постраничной разметки: содержимое отнесено к странице 1."
    )
    return DocxExtraction(document=document_json, warnings=warnings)
